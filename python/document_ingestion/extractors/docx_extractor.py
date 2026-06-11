import os
import re
from docx import Document
from models.document_schema import (DocumentSchema, Page, TextAsset, TableAsset, ImageAsset, LinkAsset)
from storage.asset_manager import AssetManager
from utils.metadata_utils import extract_basic_metadata
from utils.ocr_utils import perform_ocr
from utils.text_cleaner import TextCleaner
from docx.oxml.ns import qn
from utils.structure_detector import StructureDetector

class DOCXExtractor:

    def extract(self, file_path):
        doc = Document(file_path)
        document = DocumentSchema.create(file_name=os.path.basename(file_path), file_type="docx")
        document.metadata.update(extract_basic_metadata(file_path))
        page = Page(page_number=1)
        full_text = []

        # -------------------------
        # PARAGRAPHS + HEADINGS
        # -------------------------
        for para_idx, para in enumerate(doc.paragraphs):

            text = para.text.strip()
            if not text:
                continue

            style_name = (para.style.name if para.style else "")

            heading_level = self.detect_heading(style_name, text)

            asset = TextAsset(
                asset_id=f"para_{para_idx}",
                asset_type="text",
                text=text,
                page_number=1,
                heading_level=heading_level,
                metadata={
                    "style": style_name,
                    "source": "paragraph"
                }
            )

            page.assets.append(asset)
            full_text.append(text)

        # -------------------------
        # TABLES
        # -------------------------
        for table_idx, table in enumerate(doc.tables):

            rows = []

            for row in table.rows:
                rows.append([
                    cell.text.strip()
                    for cell in row.cells
                ])

            table_asset = TableAsset(
                asset_id=f"table_{table_idx}",
                asset_type="table",
                rows=rows,
                page_number=1,
                metadata={
                    "source": "docx_table"
                }
            )

            page.assets.append(table_asset)

        # -------------------------
        # IMAGES (DOCX PARTS)
        # -------------------------
        image_assets = self.extract_images(doc)

        page.assets.extend(image_assets)

        # -------------------------
        # LINKS (HYPERLINKS)
        # -------------------------
        link_assets = self.extract_links(doc)

        page.assets.extend(link_assets)

        document.pages.append(page)

        raw_text = "\n".join(full_text)
        document.extracted_text = TextCleaner.clean(raw_text)
        document.metadata["headings"] = (StructureDetector.detect_headings(document.extracted_text))
        return document

    # -------------------------
    # HEADING DETECTION
    # -------------------------
    def detect_heading(self, style_name, text):
        if style_name:
            match = re.search(r'heading\s+(\d+)', style_name.lower())
            if match:
                return int(match.group(1))
            

        if(len(text) < 80 and not text.endswith(".") and not text.endswith(":")):
            return 1
        return 0

    # -------------------------
    # IMAGE EXTRACTION
    # -------------------------
    def extract_images(self, doc):
        image_assets = []
        rels = doc.part._rels
        idx = 0
        for rel in rels:
            target = rels[rel].target_ref
            if "image" not in target:
                continue

            try:
                img_data = (rels[rel].target_part .blob )
                path = (AssetManager.save_binary(img_data, os.path.splitext(target)[1]) )
                ocr_text = perform_ocr(path)

                image_assets.append(
                    ImageAsset(
                        asset_id=f"img_{idx}",
                        asset_type="image",
                        image_path=path,
                        ocr_text=ocr_text,
                        metadata={
                            "source": "docx_image"
                        }
                    )
                )

                idx += 1

            except Exception as e:
                print(f"Image error: {e}")

        return image_assets

    # -------------------------
    # LINK EXTRACTION
    # -------------------------


    def extract_links(self, doc):
        links = []
        idx = 0
        try:
            for para in doc.paragraphs:
                paragraph_element = para._element
                for child in paragraph_element:
                    if not child.tag.endswith("hyperlink"):
                        continue
                    rel_id = child.get(qn("r:id"))
                    if not rel_id:
                        continue
                    try:
                        rel = doc.part.rels[rel_id]
                        url = rel.target_ref
                        text_parts = []

                        for elem in child.iter():
                            if elem.text:
                                text_parts.append(elem.text)

                        link_text = " ".join(text_parts).strip()

                        links.append(LinkAsset(
                                asset_id=f"link_{idx}",
                                asset_type="hyperlink",
                                uri=url,
                                text=link_text,
                                page_number=1,
                                metadata={ "source": "docx_link" } ) )
                        idx += 1
                    except Exception as ex:
                        print(f"Link parse error: {ex}")

        except Exception as ex:
            print( f"Hyperlink extraction error: {ex}")

        return links