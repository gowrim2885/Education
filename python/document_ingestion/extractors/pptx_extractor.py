import os
import uuid

from pptx import Presentation

from models.document_schema import (
    DocumentSchema,
    Page,
    TextAsset,
    TableAsset,
    ImageAsset
)

from storage.asset_manager import AssetManager
from utils.metadata_utils import extract_basic_metadata
from utils.text_cleaner import TextCleaner
from utils.structure_detector import StructureDetector

class PPTXExtractor:

    def extract(self, file_path):

        prs = Presentation(file_path)

        document = DocumentSchema.create(
            file_name=os.path.basename(file_path),
            file_type="pptx"
        )

        document.metadata.update(
            extract_basic_metadata(file_path)
        )

        full_text = []

        # -------------------------
        # SLIDES LOOP
        # -------------------------
        for slide_idx, slide in enumerate(prs.slides):

            page = Page(
                page_number=slide_idx + 1
            )

            slide_text = []

            # -------------------------
            # SHAPES (TEXT + OBJECTS)
            # -------------------------
            for shape_idx, shape in enumerate(slide.shapes):

                # TEXT FRAMES
                if hasattr(shape, "text") and shape.text.strip():

                    text = shape.text.strip()

                    asset = TextAsset(
                        asset_id=(
                            f"slide_{slide_idx}_"
                            f"shape_{shape_idx}"
                        ),
                        asset_type="text",
                        text=text,
                        page_number=slide_idx + 1,
                        metadata={
                            "source": "pptx_shape",
                            "shape_type": str(shape.shape_type)
                        }
                    )

                    page.assets.append(asset)
                    slide_text.append(text)

                # TABLES
                if shape.has_table:

                    table = shape.table

                    rows = []

                    for row in table.rows:
                        rows.append([
                            cell.text.strip()
                            for cell in row.cells
                        ])

                    table_asset = TableAsset(
                        asset_id=(
                            f"slide_{slide_idx}_table_{shape_idx}"
                        ),
                        asset_type="table",
                        rows=rows,
                        page_number=slide_idx + 1,
                        metadata={
                            "source": "pptx_table"
                        }
                    )

                    page.assets.append(table_asset)

                # IMAGES
                if shape.shape_type == 13:  # picture

                    image = shape.image

                    path = AssetManager.save_binary(
                        image.blob,
                        ".png"
                    )

                    image_asset = ImageAsset(
                        asset_id=(
                            f"slide_{slide_idx}_img_{shape_idx}"
                        ),
                        asset_type="image",
                        image_path=path,
                        page_number=slide_idx + 1,
                        metadata={
                            "source": "pptx_image"
                        }
                    )

                    page.assets.append(image_asset)

            # -------------------------
            # SPEAKER NOTES (VERY IMPORTANT)
            # -------------------------
            notes_text = self.extract_notes(slide)

            if notes_text:

                note_asset = TextAsset(
                    asset_id=f"slide_{slide_idx}_notes",
                    asset_type="speaker_notes",
                    text=notes_text,
                    page_number=slide_idx + 1,
                    metadata={
                        "source": "pptx_notes"
                    }
                )

                page.assets.append(note_asset)
                slide_text.append(notes_text)

            # -------------------------
            # FINALIZE SLIDE
            # -------------------------
            document.pages.append(page)

            full_text.append("\n".join(slide_text))

        raw_text = "/n".join(full_text)
        document.extracted_text = TextCleaner.clean(raw_text)
        document.metadata["headings"] = (StructureDetector.detect_headings(document.extracted_text))
        
        return document

    # -------------------------
    # SPEAKER NOTES EXTRACTION
    # -------------------------
    def extract_notes(self, slide):

        try:

            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame

            return text_frame.text.strip()

        except Exception:
            return ""