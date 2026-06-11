import uuid
import json
import os
import zipfile

from PIL import Image
from docx import Document
import pytesseract

from document_ingestion.config.settings import IMAGE_DIR

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

DOCX_PATH = r"C:\Gowri\Education-repo\python\AIProject.docx"

#Load the document
doc = Document(DOCX_PATH)

print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))

#Read Paragraphs
paragraphs = []
for para in doc.paragraphs:
    text = para.text.strip()

    if text:
        paragraphs.append(text)

# for para in paragraphs:
#     print("\n"+para)

#Read Tables
tables = []
for table in doc.tables:
    rows =[]
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    
    tables.append(rows)

# for table in tables:
#     for row in table:
#         print("\t".join(row))



# img = Image.open("C:\\Gowri\\Education-repo\\python\\Sample.png")
# text = pytesseract.image_to_string(img)
# print(text)

# print("Accuracy of OCR")
# img =img.convert("L")  # Convert to grayscale
# text = pytesseract.image_to_string(img)
# print(text)

#Read the Images
image_dir = "Extracted_Images"
os.makedirs(image_dir, exist_ok=True)
with zipfile.ZipFile(DOCX_PATH, "r") as docx:
    for file in docx.namelist():
        if file.startswith("word/media/"):
            file_name = os.path.basename(file)
            output_path = os.path.join(image_dir, file_name)
            with open(output_path, "wb") as f:
                f.write(docx.read(file))
            print("Saved", output_path)

#Create asset for each paragraph
assets = []
for ind, para in enumerate(paragraphs):
    assets.append({
        "asset_id": f"para_{ind}",
        "asset_type":"text",
        "order": ind,
        "text":para    
        })
    
#Create assets for Each Tables
for ind, table_data in enumerate(tables):
    assets.append({
        "asset_id": f"table_{ind}",
        "asset_type":"table",
        "table": table_data
    })

#Create asset for the image
for idx, image_name in enumerate(os.listdir(image_dir)):
    image_path = os.path.join(image_dir, image_name)
    img = Image.open(image_path)
    img = img.convert("L")
    # img = img.point(lambda x: 0 if x < 180 else 255,"1")
    ocr_text = pytesseract.image_to_string(img)
    assets.append({
        "asset_id": f"img_{idx}",
        "asset_type": "image",
        "image_path": image_path,
        "ocr_text": ocr_text
    })


#Create a dictionary to store all the assets
document_json ={
    "document_id":str(uuid.uuid4()),
    "file_name":os.path.basename(DOCX_PATH),
    "asset_count": len(assets), 
    "assets": assets
}

with open("document_assets.json", "w") as f:
    json.dump(document_json, f, indent=4)

print("\nJSON Saved Successfully")
print("Total Assets:", len(assets))