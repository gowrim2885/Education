import os
import io
import json
import fitz
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document
from pptx import Presentation

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'


# -------------------------
# COMMON HELPERS
# -------------------------

def detect_file_type(filename):
    return os.path.splitext(filename)[1].lower()


def ensure_dir(path):
    if not os.path.exists(path):
        os.makedirs(path)


def save_image(pix, path):
    if pix.n > 4:
        pix = fitz.Pixmap(fitz.csRGB, pix)
    pix.save(path)


def ocr_image(img):
    return pytesseract.image_to_string(img, config="--psm 6").strip()


# -------------------------
# TXT
# -------------------------
def process_txt(file):
    with open(file, "r", encoding="utf-8", errors="ignore") as f:
        return {
            "type": "txt",
            "text": f.read()
        }


# -------------------------
# DOCX
# -------------------------
def process_docx(file, output_dir):
    doc = Document(file)

    data = {
        "type": "docx",
        "paragraphs": [],
        "tables": [],
        "images": []
    }

    # paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            data["paragraphs"].append(para.text)

    # tables
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            table_data.append([cell.text.strip() for cell in row.cells])
        data["tables"].append(table_data)

    # images (raw xml extraction)
    rels = doc.part._rels
    for rel in rels:
        if "image" in rels[rel].target_ref:
            img_data = rels[rel].target_part.blob
            img_name = os.path.join(output_dir, f"doc_img_{rel}.png")
            with open(img_name, "wb") as f:
                f.write(img_data)
            data["images"].append(img_name)

    return data


# -------------------------
# PPTX
# -------------------------
def process_pptx(file, output_dir):
    prs = Presentation(file)

    data = {
        "type": "pptx",
        "slides": []
    }

    for i, slide in enumerate(prs.slides, start=1):
        slide_data = {
            "slide": i,
            "text": [],
            "images": []
        }

        for shape in slide.shapes:
            # text
            if hasattr(shape, "text") and shape.text.strip():
                slide_data["text"].append(shape.text)

            # images
            if shape.shape_type == 13:  # picture
                image = shape.image
                img_path = os.path.join(output_dir, f"slide_{i}_img.png")
                with open(img_path, "wb") as f:
                    f.write(image.blob)
                slide_data["images"].append(img_path)

        data["slides"].append(slide_data)

    return data


# -------------------------
# IMAGE
# -------------------------
def process_image(file, output_dir):
    img = Image.open(file)

    return {
        "type": "image",
        "file": file,
        "ocr_text": ocr_image(img)
    }


# -------------------------
# PDF
# -------------------------
def process_pdf(file, output_dir):
    data = {
        "type": "pdf",
        "pages": []
    }

    with fitz.open(file) as doc, pdfplumber.open(file) as pdf:

        for page_num, page in enumerate(doc, start=1):

            page_data = {
                "page": page_num,
                "text": "",
                "tables": [],
                "images": [],
                "ocr": ""
            }

            # TEXT
            text = page.get_text("text")
            if text.strip():
                page_data["text"] = text

            # TABLES
            tables = pdf.pages[page_num - 1].extract_tables()
            for table in tables:
                page_data["tables"].append(table)

            # IMAGES + OCR
            for img_index, img in enumerate(page.get_images(full=True), start=1):
                xref = img[0]
                pix = fitz.Pixmap(doc, xref)

                img_path = os.path.join(output_dir, f"page_{page_num}_img_{img_index}.png")
                save_image(pix, img_path)

                page_data["images"].append(img_path)

                # OCR on image
                pil_img = Image.open(img_path)
                ocr_text = ocr_image(pil_img)
                if ocr_text:
                    page_data["ocr"] += ocr_text + "\n"

            # FULL PAGE OCR fallback
            if not page_data["text"]:
                pix = page.get_pixmap()
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                page_data["ocr"] += ocr_image(img)

            data["pages"].append(page_data)

    return data


# -------------------------
# MAIN ROUTER
# -------------------------
def extract_file(file):

    ext = detect_file_type(file)
    output_dir = "extracted_assets"
    ensure_dir(output_dir)

    if ext == ".txt":
        return process_txt(file)

    elif ext == ".docx":
        return process_docx(file, output_dir)

    elif ext == ".pptx":
        return process_pptx(file, output_dir)

    elif ext == ".pdf":
        return process_pdf(file, output_dir)

    elif ext in [".png", ".jpg", ".jpeg"]:
        return process_image(file, output_dir)

    else:
        return {"error": "Unsupported file type"}


# -------------------------
# ENTRY POINT
# -------------------------
if __name__ == "__main__":
    file = "Sample.pdf"  # change input

    result = extract_file(file)

    # Save structured JSON
    with open("output.txt", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=4, ensure_ascii=False)

    print("✅ Extraction complete!")